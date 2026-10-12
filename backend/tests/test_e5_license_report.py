"""E5 license report: pure builders."""
import io
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

import e5_license_report as r

NOW = datetime(2026, 10, 12, 13, 0, tzinfo=timezone.utc)


def _user(name="Jane Doe", upn="jdoe@movedocs.com", enabled=True, last=None, nonint=None, synced=True):
    sia = {}
    if last is not None:
        sia["lastSignInDateTime"] = (NOW - timedelta(days=last)).isoformat()
    if nonint is not None:
        sia["lastNonInteractiveSignInDateTime"] = (NOW - timedelta(days=nonint)).isoformat()
    return {
        "displayName": name,
        "userPrincipalName": upn,
        "accountEnabled": enabled,
        "createdDateTime": "2024-01-01T00:00:00Z",
        "onPremisesSyncEnabled": synced or None,
        "signInActivity": sia or None,
    }


def test_active_enabled_user_excluded():
    assert r.classify_user(_user(last=10), NOW) is None
    assert r.classify_user(_user(last=60), NOW) is None


def test_buckets():
    assert r.classify_user(_user(enabled=False, last=5), NOW)["status"] == r.STATUS_DISABLED
    assert r.classify_user(_user(), NOW)["status"] == r.STATUS_NEVER
    assert r.classify_user(_user(last=61), NOW)["status"] == r.STATUS_IDLE_60
    assert r.classify_user(_user(last=90), NOW)["status"] == r.STATUS_IDLE_60
    assert r.classify_user(_user(last=91), NOW)["status"] == r.STATUS_IDLE_90


def test_latest_of_interactive_and_noninteractive_wins():
    row = r.classify_user(_user(last=200, nonint=75), NOW)
    assert row["days"] == 75
    assert row["status"] == r.STATUS_IDLE_60


def test_nonperson_flag_uses_word_boundaries():
    assert r.classify_user(_user(name="Conf Room 4", enabled=False), NOW)["flag"] == "Possible non-person account"
    assert r.classify_user(_user(name="Test User", enabled=False), NOW)["flag"] == "Possible non-person account"
    assert r.classify_user(_user(name="Preston Latest", enabled=False), NOW)["flag"] == ""
    assert r.classify_user(_user(name="Contessa Rhodes", enabled=False), NOW)["flag"] == ""


def test_build_rows_counts_excluded_and_sorts():
    users = [_user(last=5), _user(name="A", upn="a@x.com", enabled=False, last=100), _user(name="B", upn="b@x.com")]
    rows, excluded = r.build_rows(users, NOW)
    assert excluded == 1
    assert [x["status"] for x in rows] == [r.STATUS_DISABLED, r.STATUS_NEVER]


def test_workbook_sheets_and_summary_formulas():
    rows, excluded = r.build_rows(
        [_user(enabled=False, last=100), _user(name="B", upn="b@x.com"), _user(name="C", upn="c@x.com", last=95)], NOW
    )
    data = r.build_workbook(rows, excluded, r.E5Seats(975, 965), NOW)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "E5 Accounts"]
    acc = wb["E5 Accounts"]
    assert [c.value for c in acc[1]] == [
        "Display Name", "UPN", "Enabled", "Last Sign-In (any)", "Days Since Sign-In",
        "Created", "Synced From AD", "Status", "Flag",
    ]
    assert acc.max_row == 4
    s = wb["Summary"]
    assert s["A4"].value == "Status"
    assert s["B5"].value == "=COUNTIF('E5 Accounts'!$H$2:$H$4,A5)"
    assert s["B9"].value == "=SUM(B5:B8)"
    assert s["B11"].value == "=B5+B6+B7"
    assert s["B13"].value == 975
    assert s["B14"].value == 965


def test_summary_html_mentions_counts_and_seats():
    rows, excluded = r.build_rows([_user(enabled=False, last=100), _user(name="B", upn="b@x.com")], NOW)
    html = r.build_summary_html(rows, excluded, r.E5Seats(975, 965), NOW)
    assert "965" in html and "975" in html
    assert "Disabled" in html and "never signed in" in html
