# backend/tests/test_routes_retention.py
from __future__ import annotations

import tempfile

import pytest


@pytest.fixture(autouse=True)
def _fresh_retention_policy_store(monkeypatch):
    """retention_policy_store is a real module-level singleton also used by
    production code in routes_retention.py, and its _conn() falls back to
    the real shared Postgres database whenever DATABASE_URL is set (as it is
    in the backend container) rather than always using a private SQLite
    file. Swap in a fresh instance backed by its own temp SQLite file for
    each test instead of mutating the real singleton's data — the same
    pattern used by test_retention_policy_store.py / test_retention_cleanup_job.py
    (constructing RetentionPolicyStore(db_path=...)) and by
    test_routes_quarantine_release.py (monkeypatching the route module's
    singleton reference to a fresh instance). This also keeps state isolated
    between the test functions below that intentionally reuse the same
    team_id/channel_id, which would otherwise collide on the
    (team_id, channel_id) unique index if they shared one store."""
    from retention_policy_store import RetentionPolicyStore
    import retention_policy_store as rps_module
    import routes_retention

    store = RetentionPolicyStore(db_path=tempfile.mktemp(suffix=".db"))
    monkeypatch.setattr(rps_module, "retention_policy_store", store)
    monkeypatch.setattr(routes_retention, "retention_policy_store", store)
    return store


# Every /api/retention route is gated on the `retention` site scope as well as the
# allowlist, and the shared test_client defaults to http://testserver (which maps to
# the `primary` scope), so requests must carry the retention host explicitly or they
# correctly 404. This mirrors how the SiteContextMiddleware derives scope in
# production (host / x-forwarded-host -> get_site_scope_for_host).
RETENTION_HOST = {"host": "retention.movedocs.com"}


def _allow_retention(monkeypatch, email="test@example.com"):
    import auth
    monkeypatch.setattr(auth, "RETENTION_ALLOWED_USERS", email)


def test_connection_status_forbidden_when_not_allowlisted(test_client, monkeypatch):
    import auth
    monkeypatch.setattr(auth, "RETENTION_ALLOWED_USERS", "someone-else@example.com")
    resp = test_client.get("/api/retention/connection/status", headers=RETENTION_HOST)
    assert resp.status_code == 403


def test_connection_status_disconnected_by_default(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    resp = test_client.get("/api/retention/connection/status", headers=RETENTION_HOST)
    assert resp.status_code == 200
    assert resp.json()["status"] == "disconnected"


def test_create_policy_rejects_out_of_range_days(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    resp = test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General", "retention_days": 400},
        headers=RETENTION_HOST,
    )
    assert resp.status_code == 400


def test_create_then_confirm_policy(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    create_resp = test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General", "retention_days": 30},
        headers=RETENTION_HOST,
    )
    assert create_resp.status_code == 200
    policy = create_resp.json()
    assert policy["status"] == "pending_preview"

    confirm_resp = test_client.post(f"/api/retention/policies/{policy['id']}/confirm", headers=RETENTION_HOST)
    assert confirm_resp.status_code == 200
    assert confirm_resp.json()["status"] == "active"


def test_create_policy_conflicts_when_channel_already_has_one(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    body = {"team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General", "retention_days": 30}
    first = test_client.post("/api/retention/policies", json=body, headers=RETENTION_HOST)
    assert first.status_code == 200
    second = test_client.post("/api/retention/policies", json=body, headers=RETENTION_HOST)
    assert second.status_code == 409


def test_patch_policy_retention_days_resets_status(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    create_resp = test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General", "retention_days": 30},
        headers=RETENTION_HOST,
    )
    policy_id = create_resp.json()["id"]
    test_client.post(f"/api/retention/policies/{policy_id}/confirm", headers=RETENTION_HOST)

    patch_resp = test_client.patch(f"/api/retention/policies/{policy_id}", json={"retention_days": 60}, headers=RETENTION_HOST)
    assert patch_resp.status_code == 200
    assert patch_resp.json()["status"] == "pending_preview"
    assert patch_resp.json()["retention_days"] == 60


def test_list_runs_and_deletions_pagination(test_client, monkeypatch):
    import retention_policy_store as rps_module
    _allow_retention(monkeypatch)
    store = rps_module.retention_policy_store
    policy = store.create_policy(
        team_id="t1", team_name="Eng", channel_id="c1", channel_name="General", retention_days=30, created_by="x",
    )
    run_id = store.start_run(policy["id"])
    store.record_deletion(
        run_id=run_id, item_type="message", item_id="m1", sender_or_author="Alice",
        original_created_at="2026-01-01T00:00:00Z", status="deleted",
    )
    store.finish_run(run_id, outcome="ok", messages_deleted=1, attachments_deleted=0)

    runs_resp = test_client.get(f"/api/retention/runs?policy_id={policy['id']}", headers=RETENTION_HOST)
    assert runs_resp.status_code == 200
    assert runs_resp.json()["total"] == 1

    deletions_resp = test_client.get(f"/api/retention/deletions?run_id={run_id}", headers=RETENTION_HOST)
    assert deletions_resp.status_code == 200
    assert deletions_resp.json()["total"] == 1


def test_retention_routes_404_on_non_retention_site_scope(test_client, monkeypatch):
    """The spec gates these routes on the `retention` scope AND the allowlist.

    Without the scope gate an allowlisted operator could drive tenant-wide Teams
    deletion from it-app/azure/security/hrapp too.
    """
    _allow_retention(monkeypatch)
    for host in ["it-app.movedocs.com", "azure.movedocs.com", "security.movedocs.com", "hrapp.movedocs.com"]:
        resp = test_client.get("/api/retention/connection/status", headers={"host": host})
        assert resp.status_code == 404, host
        assert "retention.movedocs.com" in resp.json()["detail"]


def test_retention_policy_routes_404_on_non_retention_site_scope(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    azure = {"host": "azure.movedocs.com"}
    body = {"team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General", "retention_days": 30}
    assert test_client.post("/api/retention/policies", json=body, headers=azure).status_code == 404
    assert test_client.get("/api/retention/policies", headers=azure).status_code == 404
    assert test_client.get("/api/retention/teams", headers=azure).status_code == 404
    assert test_client.get("/api/retention/runs", headers=azure).status_code == 404
    assert test_client.get("/api/retention/deletions", headers=azure).status_code == 404


def test_ensure_retention_site_unit():
    import routes_retention
    from fastapi import HTTPException
    from site_context import reset_current_site_scope, set_current_site_scope

    token = set_current_site_scope("azure")
    try:
        try:
            routes_retention._ensure_retention_site()
            assert False, "expected HTTPException"
        except HTTPException as exc:
            assert exc.status_code == 404
    finally:
        reset_current_site_scope(token)

    token = set_current_site_scope("retention")
    try:
        routes_retention._ensure_retention_site()  # must not raise
    finally:
        reset_current_site_scope(token)


def _raise(exc):
    def _inner(*_args, **_kwargs):
        raise exc
    return _inner


def test_get_teams_returns_409_when_connection_is_disconnected(test_client, monkeypatch):
    import routes_retention
    from retention_graph_connection import RetentionGraphConnectionError

    _allow_retention(monkeypatch)
    monkeypatch.setattr(
        routes_retention.retention_graph_connection, "get_valid_token",
        _raise(RetentionGraphConnectionError("Retention Graph connection is not set up")),
    )

    resp = test_client.get("/api/retention/teams", headers=RETENTION_HOST)
    assert resp.status_code == 409
    assert "not set up" in resp.json()["detail"]


def test_get_teams_returns_502_on_graph_error(test_client, monkeypatch):
    import routes_retention
    from retention_graph_client import RetentionGraphError

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_teams",
        _raise(RetentionGraphError("Graph GET /teams failed: 403 Forbidden", status_code=403)),
    )

    resp = test_client.get("/api/retention/teams", headers=RETENTION_HOST)
    assert resp.status_code == 502
    assert "403" in resp.json()["detail"]


def test_get_channels_returns_409_when_connection_is_disconnected(test_client, monkeypatch):
    import routes_retention
    from retention_graph_connection import RetentionGraphConnectionError

    _allow_retention(monkeypatch)
    monkeypatch.setattr(
        routes_retention.retention_graph_connection, "get_valid_token",
        _raise(RetentionGraphConnectionError("Retention Graph connection is disconnected")),
    )

    resp = test_client.get("/api/retention/teams/t1/channels", headers=RETENTION_HOST)
    assert resp.status_code == 409


def test_get_channels_returns_502_on_graph_error(test_client, monkeypatch):
    import routes_retention
    from retention_graph_client import RetentionGraphError

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_channels", _raise(RetentionGraphError("throttled", status_code=429)),
    )

    resp = test_client.get("/api/retention/teams/t1/channels", headers=RETENTION_HOST)
    assert resp.status_code == 502


def test_get_teams_filters_by_case_insensitive_name_substring(test_client, monkeypatch):
    import routes_retention

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_teams",
        lambda _token: [
            {"id": "t1", "name": "Engineering"},
            {"id": "t2", "name": "Sales"},
            {"id": "t3", "name": "engineering leadership"},
        ],
    )
    monkeypatch.setattr(routes_retention.retention_directory_cache, "channels_by_team_snapshot", lambda: {})

    resp = test_client.get("/api/retention/teams", params={"q": "ENGINEER"}, headers=RETENTION_HOST)
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 2
    assert {t["id"] for t in body["items"]} == {"t1", "t3"}


def test_get_teams_matches_on_a_channel_name_via_the_directory_cache(test_client, monkeypatch):
    # Graph has no "search channels across every team" endpoint, so a team with no
    # name match but a matching channel must still surface via the cached directory.
    import routes_retention

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_teams",
        lambda _token: [{"id": "t1", "name": "Libra Production Support"}, {"id": "t2", "name": "Sales"}],
    )
    monkeypatch.setattr(
        routes_retention.retention_directory_cache, "channels_by_team_snapshot",
        lambda: {
            "t1": [{"id": "c1", "name": "Incident Response"}, {"id": "c2", "name": "General"}],
            "t2": [{"id": "c3", "name": "General"}],
        },
    )

    resp = test_client.get("/api/retention/teams", params={"q": "incident"}, headers=RETENTION_HOST)
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    assert body["items"][0]["id"] == "t1"


def test_get_teams_channel_match_tolerates_a_team_missing_from_the_cache(test_client, monkeypatch):
    # A team created since the cache's last refresh (or before it has ever run) is
    # absent from the snapshot entirely — that must not error, just not match on
    # channel name for that one team until the next refresh.
    import routes_retention

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_teams",
        lambda _token: [{"id": "t1", "name": "Brand New Team"}],
    )
    monkeypatch.setattr(routes_retention.retention_directory_cache, "channels_by_team_snapshot", lambda: {})

    resp = test_client.get("/api/retention/teams", params={"q": "incident"}, headers=RETENTION_HOST)
    assert resp.status_code == 200
    assert resp.json()["total"] == 0


def test_get_teams_does_not_read_the_directory_cache_when_search_is_empty(test_client, monkeypatch):
    import routes_retention

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(routes_retention, "_list_teams", lambda _token: [{"id": "t1", "name": "Sales"}])
    snapshot_calls: list[None] = []
    monkeypatch.setattr(
        routes_retention.retention_directory_cache, "channels_by_team_snapshot",
        lambda: snapshot_calls.append(None) or {},
    )

    resp = test_client.get("/api/retention/teams", headers=RETENTION_HOST)
    assert resp.status_code == 200
    assert snapshot_calls == []


def test_get_teams_search_composes_with_pagination(test_client, monkeypatch):
    # Filtering must happen before offset/limit slicing, or `total` and the
    # returned page would reflect the unfiltered list instead of the search result.
    import routes_retention

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_teams",
        lambda _token: [
            {"id": "t1", "name": "Engineering Alpha"},
            {"id": "t2", "name": "Sales"},
            {"id": "t3", "name": "Engineering Beta"},
            {"id": "t4", "name": "Engineering Gamma"},
        ],
    )
    monkeypatch.setattr(routes_retention.retention_directory_cache, "channels_by_team_snapshot", lambda: {})

    resp = test_client.get(
        "/api/retention/teams", params={"q": "engineering", "limit": 2, "offset": 1}, headers=RETENTION_HOST,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 3
    assert [t["id"] for t in body["items"]] == ["t3", "t4"]


def test_get_channels_filters_by_case_insensitive_name_substring(test_client, monkeypatch):
    import routes_retention

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(
        routes_retention, "_list_channels",
        lambda _token, _team_id: [
            {"id": "c1", "name": "General"},
            {"id": "c2", "name": "Random"},
            {"id": "c3", "name": "general-archive"},
        ],
    )

    resp = test_client.get("/api/retention/teams/t1/channels", params={"q": "general"}, headers=RETENTION_HOST)
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 2
    assert {c["id"] for c in body["items"]} == {"c1", "c3"}


def test_teams_and_channels_routes_are_sync_defs_not_async():
    """Blocking `requests` Graph enumeration must not run on the shared event loop.

    Declaring these handlers `async def` would stall every other host's requests
    for the duration of the Graph calls; plain `def` puts them in FastAPI's
    threadpool, matching routes_ad.py / routes_tools.py.
    """
    import inspect
    import routes_retention

    assert not inspect.iscoroutinefunction(routes_retention.get_teams)
    assert not inspect.iscoroutinefunction(routes_retention.get_channels)


def test_preview_returns_502_on_graph_error(test_client, monkeypatch):
    import routes_retention
    from retention_graph_client import RetentionGraphError

    _allow_retention(monkeypatch)
    create_resp = test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General", "retention_days": 30},
        headers=RETENTION_HOST,
    )
    policy_id = create_resp.json()["id"]

    async def _boom(_policy_id):
        raise RetentionGraphError("Graph GET messages failed: 403 Forbidden", status_code=403)

    monkeypatch.setattr(routes_retention.retention_cleanup_job, "compute_preview", _boom)

    resp = test_client.get(f"/api/retention/policies/{policy_id}/preview", headers=RETENTION_HOST)
    assert resp.status_code == 502


def _xlsx_bytes(headers, rows):
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_channels_returns_xlsx_with_known_channels_and_policy_state(test_client, monkeypatch):
    import routes_retention
    from openpyxl import load_workbook
    import io

    _allow_retention(monkeypatch)
    monkeypatch.setattr(routes_retention.retention_graph_connection, "get_valid_token", lambda: "token")
    monkeypatch.setattr(routes_retention, "_list_teams", lambda _token: [{"id": "t1", "name": "Engineering"}])
    monkeypatch.setattr(
        routes_retention.retention_directory_cache, "channels_by_team_snapshot",
        lambda: {"t1": [{"id": "c1", "name": "General"}]},
    )
    test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Engineering", "channel_id": "c1", "channel_name": "General", "retention_days": 30},
        headers=RETENTION_HOST,
    )

    resp = test_client.get("/api/retention/export", headers=RETENTION_HOST)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == ["team_id", "team_name", "channel_id", "channel_name", "current_status", "current_retention_days", "retention_days"]
    data_row = [c.value for c in ws[2]]
    assert data_row == ["t1", "Engineering", "c1", "General", "pending_preview", 30, 30]


def test_export_channels_returns_409_when_disconnected(test_client, monkeypatch):
    import routes_retention
    from retention_graph_connection import RetentionGraphConnectionError

    _allow_retention(monkeypatch)
    monkeypatch.setattr(
        routes_retention.retention_graph_connection, "get_valid_token",
        _raise(RetentionGraphConnectionError("not connected")),
    )
    resp = test_client.get("/api/retention/export", headers=RETENTION_HOST)
    assert resp.status_code == 409


def test_import_preview_rejects_a_non_xlsx_filename(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    resp = test_client.post(
        "/api/retention/import/preview",
        files={"file": ("channels.csv", b"team_id,channel_id,retention_days", "text/csv")},
        headers=RETENTION_HOST,
    )
    assert resp.status_code == 400
    assert "xlsx" in resp.json()["detail"].lower()


def test_import_preview_rejects_a_workbook_missing_required_columns(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    content = _xlsx_bytes(["team_id", "channel_id"], [["t1", "c1"]])
    resp = test_client.post(
        "/api/retention/import/preview",
        files={"file": ("channels.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=RETENTION_HOST,
    )
    assert resp.status_code == 400
    assert "retention_days" in resp.json()["detail"]


def test_import_preview_classifies_create_update_skip_and_error_rows(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Existing", "channel_id": "c1", "channel_name": "General", "retention_days": 30},
        headers=RETENTION_HOST,
    )
    content = _xlsx_bytes(
        ["team_id", "team_name", "channel_id", "channel_name", "retention_days"],
        [
            ["t1", "Existing", "c1", "General", 45],       # update
            ["t2", "New Team", "c2", "Random", 60],        # create
            ["t3", "Another Team", "c3", "Standup", ""],   # skip: blank retention_days
            ["t4", "Bad Team", "c4", "Bad", 400],           # error: out of range
        ],
    )
    resp = test_client.post(
        "/api/retention/import/preview",
        files={"file": ("channels.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=RETENTION_HOST,
    )
    assert resp.status_code == 200
    rows = {r["channel_id"]: r for r in resp.json()["rows"]}
    assert rows["c1"]["action"] == "update"
    assert rows["c1"]["current_status"] == "pending_preview"
    assert rows["c2"]["action"] == "create"
    assert rows["c3"]["action"] == "skip"
    assert rows["c4"]["action"] == "error"
    assert "between 1 and 365" in rows["c4"]["error"]


def test_import_preview_flags_duplicate_rows_for_the_same_channel(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    content = _xlsx_bytes(
        ["team_id", "team_name", "channel_id", "channel_name", "retention_days"],
        [["t1", "Eng", "c1", "General", 30], ["t1", "Eng", "c1", "General", 60]],
    )
    resp = test_client.post(
        "/api/retention/import/preview",
        files={"file": ("channels.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=RETENTION_HOST,
    )
    rows = resp.json()["rows"]
    assert rows[0]["action"] == "create"
    assert rows[1]["action"] == "error"
    assert "Duplicate" in rows[1]["error"]


def test_import_apply_creates_and_updates_policies(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    test_client.post(
        "/api/retention/policies",
        json={"team_id": "t1", "team_name": "Existing", "channel_id": "c1", "channel_name": "General", "retention_days": 30},
        headers=RETENTION_HOST,
    )
    content = _xlsx_bytes(
        ["team_id", "team_name", "channel_id", "channel_name", "retention_days"],
        [["t1", "Existing", "c1", "General", 45], ["t2", "New Team", "c2", "Random", 60]],
    )
    resp = test_client.post(
        "/api/retention/import/apply",
        files={"file": ("channels.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=RETENTION_HOST,
    )
    assert resp.status_code == 200
    results = {r["channel_id"]: r for r in resp.json()["rows"]}
    assert results["c1"]["result"] == "updated"
    assert results["c2"]["result"] == "created"

    from retention_policy_store import retention_policy_store as store
    updated = store.get_policy_for_channel("t1", "c1")
    assert updated["retention_days"] == 45
    assert updated["status"] == "pending_preview"
    created = store.get_policy_for_channel("t2", "c2")
    assert created["retention_days"] == 60
    assert created["status"] == "pending_preview"


def test_import_apply_does_not_touch_skip_or_error_rows(test_client, monkeypatch):
    _allow_retention(monkeypatch)
    content = _xlsx_bytes(
        ["team_id", "team_name", "channel_id", "channel_name", "retention_days"],
        [["t1", "Eng", "c1", "General", ""], ["t2", "Bad", "c2", "Bad", 9999]],
    )
    resp = test_client.post(
        "/api/retention/import/apply",
        files={"file": ("channels.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=RETENTION_HOST,
    )
    results = {r["channel_id"]: r for r in resp.json()["rows"]}
    assert results["c1"]["result"] == "skipped"
    assert results["c2"]["result"] == "skipped"

    from retention_policy_store import retention_policy_store as store
    assert store.get_policy_for_channel("t1", "c1") is None
    assert store.get_policy_for_channel("t2", "c2") is None
