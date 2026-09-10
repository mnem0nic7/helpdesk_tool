# backend/tests/test_retention_import_export.py
from __future__ import annotations

import io


def _workbook_bytes(headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_build_export_workbook_writes_header_and_rows():
    from retention_import_export import EXPORT_COLUMNS, build_export_workbook
    rows = [{
        "team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General",
        "current_status": "active", "current_retention_days": 30, "retention_days": 30,
    }]
    wb = build_export_workbook(rows)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == EXPORT_COLUMNS
    data_row = [cell.value for cell in ws[2]]
    assert data_row == ["t1", "Eng", "c1", "General", "active", 30, 30]


def test_parse_import_workbook_reads_columns_by_name_regardless_of_order():
    content = _workbook_bytes(
        ["retention_days", "channel_id", "team_id", "channel_name", "team_name"],
        [[45, "c1", "t1", "General", "Eng"]],
    )
    from retention_import_export import parse_import_workbook
    rows = parse_import_workbook(content)
    assert rows == [{
        "team_id": "t1", "team_name": "Eng", "channel_id": "c1", "channel_name": "General",
        "retention_days": "45",
    }]


def test_parse_import_workbook_normalizes_whole_number_floats():
    # openpyxl reads a numeric cell typed "45" back as the float 45.0 — this must not
    # leak into the raw string as "45.0" for the caller's later int() parsing.
    content = _workbook_bytes(["team_id", "channel_id", "retention_days"], [["t1", "c1", 45.0]])
    from retention_import_export import parse_import_workbook
    rows = parse_import_workbook(content)
    assert rows[0]["retention_days"] == "45"


def test_parse_import_workbook_skips_fully_blank_rows():
    content = _workbook_bytes(["team_id", "channel_id", "retention_days"], [["t1", "c1", 30], [None, None, None]])
    from retention_import_export import parse_import_workbook
    rows = parse_import_workbook(content)
    assert len(rows) == 1


def test_parse_import_workbook_raises_on_missing_required_column():
    content = _workbook_bytes(["team_id", "channel_id"], [["t1", "c1"]])
    from retention_import_export import parse_import_workbook
    try:
        parse_import_workbook(content)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "retention_days" in str(exc)


def test_parse_import_workbook_raises_on_unreadable_content():
    from retention_import_export import parse_import_workbook
    try:
        parse_import_workbook(b"not a real xlsx file")
        assert False, "expected ValueError"
    except ValueError:
        pass
